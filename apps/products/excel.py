"""
MERRIGE ERP - Барааны Excel Import/Export
Import: Excel файлаас барааг мөр мөрөөр нь ЯГ ижил
ProductCreateUpdateSerializer + ProductService.create_product()-оор
шалгаж/үүсгэнэ — ингэснээр Excel-ээр орж ирсэн өгөгдөл ч API-тай яг
ижил бизнес дүрэм (код давхардахгүй, үнэ 0-ээс их байх ёстой) дагана.
Export: Одоогийн барааны каталогийг Excel болгож гаргана.
"""

from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from apps.products.models import ProductStatus
from apps.products.serializers import ProductCreateUpdateSerializer
from apps.products.services import ProductService
from apps.shared.exceptions import ХүчинтэйБайдлынАлдаа

IMPORT_HEADERS = ["Код", "Нэр", "Тайлбар", "Материал", "Арчилгааны заавар", "Үнэ", "Төлөв"]
EXPORT_HEADERS = [
    "Код",
    "Нэр",
    "Тайлбар",
    "Материал",
    "Арчилгааны заавар",
    "Үнэ",
    "Төлөв",
    "Үүсгэсэн огноо",
]

_STATUS_LABEL_TO_VALUE = {label: value for value, label in ProductStatus.choices}


def import_products_from_excel(file_obj):
    """Excel файлыг мөр мөрөөр нь боловсруулж, [{row_number, code, success,
    errors}, ...] хэлбэрээр дэлгэрэнгүй үр дүн буцаана. Мөр бүрийг
    тусдаа шалгаж, зөв мөрийг үүсгэнэ, буруу мөрийг алгасна (бусад
    зөв мөрөнд нөлөөлөхгүй)."""
    try:
        workbook = load_workbook(file_obj, data_only=True)
    except Exception as exc:
        raise ХүчинтэйБайдлынАлдаа(
            "Excel файлыг уншиж чадсангүй. Файл эвдэрсэн эсвэл .xlsx форматгүй байна."
        ) from exc

    sheet = workbook.active
    results = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue

        padded_row = list(row) + [None] * (7 - len(row))
        code, name, description, material, care_instructions, price, status_label = padded_row[:7]

        errors = {}
        price_value = None
        if price in (None, ""):
            errors["price"] = ["Үнэ заавал шаардлагатай."]
        else:
            try:
                price_value = Decimal(str(price))
            except InvalidOperation:
                errors["price"] = ["Үнэ тоон утга байх ёстой."]

        status_value = ProductStatus.DRAFT
        if status_label not in (None, ""):
            status_value = _STATUS_LABEL_TO_VALUE.get(str(status_label).strip())
            if status_value is None:
                errors["status"] = [
                    "Төлөв буруу байна. Зөвшөөрөгдсөн утга: "
                    + ", ".join(_STATUS_LABEL_TO_VALUE)
                ]

        if errors:
            results.append(
                {"row_number": row_number, "code": code, "success": False, "errors": errors}
            )
            continue

        serializer = ProductCreateUpdateSerializer(
            data={
                "code": str(code).strip() if code else "",
                "name": str(name).strip() if name else "",
                "description": description or "",
                "material": material or "",
                "care_instructions": care_instructions or "",
                "price": price_value,
                "status": status_value,
            }
        )
        if not serializer.is_valid():
            results.append(
                {
                    "row_number": row_number,
                    "code": code,
                    "success": False,
                    "errors": serializer.errors,
                }
            )
            continue

        ProductService.create_product(serializer.validated_data)
        results.append({"row_number": row_number, "code": code, "success": True, "errors": None})

    return results


def build_products_export_rows(products):
    return [
        [
            product.code,
            product.name,
            product.description,
            product.material,
            product.care_instructions,
            float(product.price),
            product.get_status_display(),
            product.created_at.strftime("%Y-%m-%d %H:%M"),
        ]
        for product in products
    ]
