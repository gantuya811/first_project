"""
MERRIGE ERP - Excel/PDF экспортын нийтлэг туслах функцууд
Анх apps.reports дотор байсан (STEP11) боловч apps.products (Excel Import/
Export, STEP15), apps.orders (Нэгдсэн захиалгын Excel экспорт) зэрэг олон
app ашиглах ёстой цэвэр infra тул apps.shared-руу зөв байрлуулав.

PDF-ийн тухай тэмдэглэл: reportlab-ийн стандарт Helvetica фонт Кирилл
(Монгол) үсэг ДЭМЖДЭГГҮЙ. Иймд "apps/shared/fonts/DejaVuSans.ttf" (болон
Bold хувилбар)-г төсөлд шууд bundle хийсэн. DejaVu Sans бол чөлөөт эрх
бүхий (Bitstream Vera License-ийн уламжлал, дэлгэрэнгүйг
apps/shared/fonts/DEJAVU_LICENSE.txt-c үзнэ үү), Cyrillic-ийг бүрэн
дэмждэг фонт. Зөвхөн энэ .ttf файлыг ашигласан тул matplotlib зэрэг том
сан руу хамаарал үүсгээгүй болно.
"""

import io
import os

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_FONTS_DIR = os.path.join(os.path.dirname(__file__), "fonts")
_FONT_REGULAR = "DejaVuSans"
_FONT_BOLD = "DejaVuSans-Bold"
_fonts_registered = False


def _ensure_fonts_registered():
    global _fonts_registered
    if _fonts_registered:
        return
    pdfmetrics.registerFont(
        TTFont(_FONT_REGULAR, os.path.join(_FONTS_DIR, "DejaVuSans.ttf"))
    )
    pdfmetrics.registerFont(
        TTFont(_FONT_BOLD, os.path.join(_FONTS_DIR, "DejaVuSans-Bold.ttf"))
    )
    _fonts_registered = True


def export_rows_to_excel(filename, headers, rows, sheet_title="Тайлан"):
    """headers: [str, ...]; rows: [[утга, ...], ...]. Excel файлыг HTTP
    хариу болгон буцаана (Content-Disposition: attachment)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]  # Excel sheet нэр 31 тэмдэгтээс ихгүй

    sheet.append(headers)
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        sheet.append(row)

    for column_cells in sheet.columns:
        max_length = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in column_cells
        )
        column_letter = column_cells[0].column_letter
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 50)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_rows_to_pdf(filename, title, headers, rows):
    """headers: [str, ...]; rows: [[утга, ...], ...]. PDF файлыг HTTP
    хариу болгон буцаана. Монгол (Cyrillic) текстийг DejaVu Sans фонтоор
    зөв дүрслэнэ."""
    _ensure_fonts_registered()

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title=title,
    )

    title_style = ParagraphStyle(
        name="Гарчиг", fontName=_FONT_BOLD, fontSize=16, leading=20
    )
    elements = [Paragraph(title, title_style), Spacer(1, 14)]

    table_data = [headers] + [[str(value) for value in row] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), _FONT_BOLD),
                ("FONTNAME", (0, 1), (-1, -1), _FONT_REGULAR),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elements.append(table)
    document.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
