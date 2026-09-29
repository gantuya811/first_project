"""
MERRIGE ERP - Excel Import/Export (STEP15) тестүүд
"""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from apps.products.excel import import_products_from_excel
from apps.products.models import Product

pytestmark = pytest.mark.django_db


def _build_import_file(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Код", "Нэр", "Тайлбар", "Материал", "Арчилгааны заавар", "Үнэ", "Төлөв"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


class TestProductImport:
    def test_valid_row_creates_product(self):
        file_obj = _build_import_file([["IMP01", "Тест", "", "", "", 10000, "Идэвхтэй"]])
        results = import_products_from_excel(file_obj)
        assert results[0]["success"] is True
        assert Product.objects.filter(code="IMP01").exists()

    def test_invalid_price_rejected(self):
        file_obj = _build_import_file([["IMP02", "Тест", "", "", "", "арван мянга", ""]])
        results = import_products_from_excel(file_obj)
        assert results[0]["success"] is False
        assert "price" in results[0]["errors"]

    def test_invalid_status_label_rejected(self):
        file_obj = _build_import_file([["IMP03", "Тест", "", "", "", 5000, "БАЙХГҮЙ"]])
        results = import_products_from_excel(file_obj)
        assert results[0]["success"] is False
        assert "status" in results[0]["errors"]

    def test_duplicate_code_rejected_same_as_api(self, product):
        file_obj = _build_import_file([[product.code, "Х", "", "", "", 1000, ""]])
        results = import_products_from_excel(file_obj)
        assert results[0]["success"] is False
        assert "бүртгэлтэй" in str(results[0]["errors"])

    def test_mixed_valid_invalid_rows_independent(self):
        file_obj = _build_import_file(
            [
                ["OK01", "Зөв", "", "", "", 1000, ""],
                ["", "Кодгүй", "", "", "", 1000, ""],
                ["OK02", "Зөв 2", "", "", "", 2000, ""],
            ]
        )
        results = import_products_from_excel(file_obj)
        assert [r["success"] for r in results] == [True, False, True]
        assert Product.objects.filter(code="OK01").exists()
        assert Product.objects.filter(code="OK02").exists()

    def test_blank_rows_skipped(self):
        file_obj = _build_import_file(
            [["OK03", "Зөв", "", "", "", 1000, ""], [None, None, None, None, None, None, None]]
        )
        results = import_products_from_excel(file_obj)
        assert len(results) == 1


class TestProductExportAPI:
    def test_export_contains_created_products(self, admin_client, product):
        response = admin_client.get("/api/v1/products/export/")
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        codes = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert product.code in codes


class TestProductImportAPI:
    def test_seller_cannot_import(self, main_seller_client):
        file_obj = _build_import_file([["X", "X", "", "", "", 1000, ""]])
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            "products.xlsx", file_obj.read(), content_type="application/vnd.ms-excel"
        )
        response = main_seller_client.post(
            "/api/v1/products/import/", {"file": upload}, format="multipart"
        )
        assert response.status_code == 403

    def test_non_xlsx_extension_rejected(self, admin_client):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile("data.csv", b"code,name\n", content_type="text/csv")
        response = admin_client.post(
            "/api/v1/products/import/", {"file": upload}, format="multipart"
        )
        assert response.status_code == 400
